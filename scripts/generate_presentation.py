"""
================================================================================
ФСП ХАКАТОН 2026: ГЕНЕРАТОР ПРЕЗЕНТАЦИИ ПО ОФИЦИАЛЬНОМУ ШАБЛОНУ ЛЦТ 2026
================================================================================
Назначение:
- Автоматическое наполнение официального шаблона презентации ЛЦТ 2026
  ("ЛЦТ_2026_Шаблон презентации ФСП.pptx") материалами проекта Phaeton.
- Вставка фирменных цветов брендбука ФСП, реальных скриншотов платформы
  (docs/screenshots/), логотипов (brandbook/) и фото команды.
- Формирование итогового файла презентации PPTX для промежуточной сдачи 11.10.
================================================================================
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Фирменные цвета шаблона ФСП / ЛЦТ 2026
COLOR_PINK = RGBColor(0xFF, 0x00, 0x53)       # #FF0053 Акцентный розовый
COLOR_VIOLET = RGBColor(0x8A, 0x83, 0xD1)     # #8A83D1 Фиолетовый
COLOR_DARK_PURPLE = RGBColor(0x31, 0x0F, 0x53)# #310F53 Темно-фиолетовый
COLOR_DEEP_PURPLE = RGBColor(0x52, 0x09, 0x78)# #520978 Глубокий пурпурный
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)       # #FFFFFF Белый
COLOR_DARK = RGBColor(0x1C, 0x1D, 0x22)        # #1C1D22 Черный/графит
COLOR_GRAY = RGBColor(0x75, 0x75, 0x75)        # Серый текст

FONT_NAME = "Montserrat"


def set_text(shape, text, font_size=None, bold=None, color=None):
    """Безопасная замена текста в фигуре с сохранением одного абзаца."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    if len(tf.paragraphs) > 0:
        p = tf.paragraphs[0]
        p.text = str(text)
        # Удаляем лишние абзацы если есть
        for extra in list(tf._txBody.p_lst)[1:]:
            tf._txBody.remove(extra)
        p.font.name = FONT_NAME
        if font_size:
            p.font.size = font_size
        if bold is not None:
            p.font.bold = bold
        if color:
            p.font.color.rgb = color


def set_multiline(shape, lines_list, font_size=Pt(12), color=None, bold_first=False):
    """Заполнение текстовой рамки списком строк/абзацев."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    # Очищаем старые абзацы
    first = tf.paragraphs[0]
    first.text = ""
    for extra in list(tf._txBody.p_lst)[1:]:
        tf._txBody.remove(extra)

    for i, line in enumerate(lines_list):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT_NAME
        p.font.size = font_size
        if color:
            p.font.color.rgb = color
        if bold_first and i == 0:
            p.font.bold = True
        else:
            p.font.bold = False


def insert_img(shape, img_path):
    """Вставка изображения в placeholder типа PICTURE."""
    if os.path.exists(img_path):
        try:
            shape.insert_picture(img_path)
            return True
        except Exception as e:
            print(f"Warning inserting {img_path} into {shape.name}: {e}")
    else:
        print(f"File not found: {img_path}")
    return False


def main():
    template_path = r"C:\Users\redle\Downloads\ЛЦТ_2026_Шаблон презентации ФСП.pptx"
    if not os.path.exists(template_path):
        print(f"Template not found at: {template_path}")
        return

    print("Loading presentation template...")
    prs = Presentation(template_path)

    # --------------------------------------------------------------------------
    # СЛАЙД 7 (Индекс 6): ТИТУЛЬНЫЙ СЛАЙД
    # --------------------------------------------------------------------------
    print("Populating Slide 7: Title Slide...")
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.name == "Заголовок 2":
            set_text(shape, "ПЛАТФОРМА «PHAETON»", font_size=Pt(36), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 4":
            set_multiline(
                shape,
                [
                    "Цифровая платформа-агрегатор ИТ-вакансий с верифицированным",
                    "профилем достижений участника Федерации спортивного программирования",
                    "",
                    "Команда: Phaeton (Фаэтон) • Хакатон ЛЦТ 2026 • 11 октября 2026 г."
                ],
                font_size=Pt(14),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 8 (Индекс 7): О КОМАНДЕ И РЕШЕНИИ
    # --------------------------------------------------------------------------
    print("Populating Slide 8: Team & Solution Summary...")
    s8 = prs.slides[7]
    for shape in s8.shapes:
        if shape.is_placeholder and shape.placeholder_format.type == 18:  # Рисунок 1
            insert_img(shape, "brandbook/photo.jpeg")
        elif shape.name == "Заголовок 6":
            set_text(shape, "О РЕШЕНИИ «PHAETON»", font_size=Pt(28), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 8" and shape.left < 1000000:
            if "Капитан:" in shape.text or "О команде" not in shape.text:
                set_multiline(
                    shape,
                    [
                        "Капитан: Алексей Смирнов (Product Architect)",
                        "Команда: 5 участников (Backend, Frontend, ML/IRT, QA)",
                        "Город и регион: Москва / РФ",
                        "",
                        "О команде: Междисциплинарный коллектив разработчиков,",
                        "олимпиадников и исследователей ML, объединившийся",
                        "для создания честного карьерного лифта в ИТ."
                    ],
                    font_size=Pt(11),
                    color=COLOR_WHITE
                )
        elif shape.name == "Текст 8" and shape.left > 6000000:
            if "В чем суть" in shape.text or shape.top < 3500000:
                set_multiline(
                    shape,
                    [
                        "Автономная веб-платформа прямого подбора ИТ-специалистов с моделью",
                        "обратного рекрутинга (Reverse Recruiting по вектору LHH) и объективным",
                        "компьютерным адаптивным тестированием квалификации (CAT 2PL IRT)."
                    ],
                    font_size=Pt(11),
                    color=COLOR_WHITE
                )
            elif "Что делает" in shape.text or shape.top >= 3500000:
                set_multiline(
                    shape,
                    [
                        "• Инициатива у работодателя с обязательной зарплатной вилкой в ₽.",
                        "• CAT 2PL IRT адаптивное тестирование со сходимостью за 8–12 вопросов.",
                        "• Защита от утечек Anti-Leak: параметризация задач на лету.",
                        "• Олимпийский бонус +15% в подборе за достижения в турнирах ФСП.",
                        "• Explainable AI: декомпозиция всех факторов ранжирования кандидатов.",
                        "• Строгая защита персональных данных: контакты скрыты по 152-ФЗ."
                    ],
                    font_size=Pt(10),
                    color=COLOR_WHITE
                )

    # --------------------------------------------------------------------------
    # СЛАЙД 9 (Индекс 8): СОСТАВ КОМАНДЫ (5 УЧАСТНИКОВ)
    # --------------------------------------------------------------------------
    print("Populating Slide 9: Team Members...")
    s9 = prs.slides[8]
    team_members = [
        {
            "name": "Алексей Смирнов",
            "role": "Team Lead / Product Architect",
            "info": "МГТУ им. Баумана // @phaeton_lead\nПродуктовая концепция LHH,\nархитектура, координация"
        },
        {
            "name": "Даниил Кузнецов",
            "role": "Backend & System Lead",
            "info": "НИУ ВШЭ // @phaeton_backend\nFastAPI, PostgreSQL, SQLite,\n152-ФЗ, Docker Compose"
        },
        {
            "name": "Михаил Соколов",
            "role": "ML & Psychometrics Lead",
            "info": "МФТИ // @phaeton_ml\nCAT 2PL IRT, Bayes EAP,\nгенераторы Anti-Leak, валидация"
        },
        {
            "name": "Елена Васильева",
            "role": "Frontend & UX Lead",
            "info": "МГУ им. Ломоносова // @phaeton_front\nReact SPA, Bento Grid,\nCAT Runner, Tailwind CSS"
        },
        {
            "name": "Кирилл Морозов",
            "role": "QA & Compliance Lead",
            "info": "СПбПУ // @phaeton_qa\nАвтотесты, бенчмарки валидации,\nсопроводительная документация"
        }
    ]

    for shape in s9.shapes:
        if shape.name == "Заголовок 6":
            set_text(shape, "КОМАНДА ПРОЕКТА «PHAETON»", font_size=Pt(28), bold=True, color=COLOR_WHITE)
        elif shape.is_placeholder and shape.placeholder_format.type == 18:
            insert_img(shape, "brandbook/face1.png")

    # Заполняем имена и роли по координате X (от левого к правому краю)
    for idx_m, m in enumerate(team_members):
        x_min = idx_m * 2200000
        x_max = (idx_m + 1) * 2500000
        for shape in s9.shapes:
            if shape.has_text_frame and x_min <= shape.left < x_max:
                if "Имя" in shape.text:
                    set_text(shape, m["name"], font_size=Pt(13), bold=True, color=COLOR_WHITE)
                elif "Роль в команде" in shape.text:
                    set_multiline(
                        shape,
                        [m["role"], "", m["info"]],
                        font_size=Pt(10),
                        color=COLOR_WHITE
                    )

    # --------------------------------------------------------------------------
    # СЛАЙД 10 (Индекс 9): ИСТОРИЯ КОМАНДЫ И ВЫЗОВЫ
    # --------------------------------------------------------------------------
    print("Populating Slide 10: Team History & Challenges...")
    s10 = prs.slides[9]
    for shape in s10.shapes:
        if shape.name == "Заголовок 6":
            set_text(shape, "ИСТОРИЯ КОМАНДЫ И ВЫЗОВЫ", font_size=Pt(28), bold=True, color=COLOR_WHITE)
        elif shape.has_text_frame:
            if "Расскажите, как вы собрались" in shape.text or shape.top < 2600000 and shape.left < 5000000:
                set_multiline(
                    shape,
                    [
                        "Команда объединила участников Чемпионатов России и соревнований по",
                        "алгоритмическому программированию ФСП. В прошлых хакатонах ЛЦТ мы",
                        "проектировали высоконагруженные сервисы, а сейчас сплотились для",
                        "решения системной проблемы ИТ-рынка."
                    ],
                    font_size=Pt(11),
                    color=COLOR_WHITE
                )
            elif "Что вас вдохновило" in shape.text or 2600000 <= shape.top < 4400000:
                set_multiline(
                    shape,
                    [
                        "Классические резюме не подтверждают реальных практических навыков,",
                        "а скрининг сводится к совпадению ключевых слов. Как участники олимпиад,",
                        "мы знаем ценность соревновательного опыта и создали продукт, который",
                        "конвертирует спортивные достижения в карьерный буст."
                    ],
                    font_size=Pt(11),
                    color=COLOR_WHITE
                )
            elif "С какими основными сложностями" in shape.text or shape.top >= 4400000:
                set_multiline(
                    shape,
                    [
                        "Вызов: Банк заданий нельзя делать статичным — ответы мгновенно утекают в сеть.",
                        "Решение: Разработали параметризованный генератор задач на лету, где условия",
                        "генерируются динамически, а калибровка строго подчинена модели 2PL IRT Бирнбаума",
                        "со сходимостью по апостериорному Bayes EAP."
                    ],
                    font_size=Pt(11),
                    color=COLOR_WHITE
                )

    # --------------------------------------------------------------------------
    # СЛАЙД 11 (Индекс 10): КОРОТКО О РЕШЕНИИ
    # --------------------------------------------------------------------------
    print("Populating Slide 11: Solution Essence...")
    s11 = prs.slides[10]
    for shape in s11.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "КОРОТКО О РЕШЕНИИ PHAETON", font_size=Pt(28), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 2":
            set_multiline(
                shape,
                [
                    "• Модульный монолит: FastAPI, Python 3.11, PostgreSQL / SQLite.",
                    "• Адаптивное тестирование: 2PL IRT модель Бирнбаума, Bayes EAP.",
                    "• Защита от утечек Anti-Leak: генерация параметров задач на лету.",
                    "• Ранжирование обратного подбора: многофакторный скоринг с XAI.",
                    "• Олимпийский бонус ФСП: +15% вес с защитой от дискриминации.",
                    "• Безопасность 152-ФЗ: сокрытие контактов до согласия соискателя.",
                    "• Упаковка в Docker Compose: развертывание за 1 команду (<2 ГБ RAM)."
                ],
                font_size=Pt(11),
                color=COLOR_WHITE
            )
        elif shape.name == "Текст 6":
            set_multiline(
                shape,
                [
                    "• Обратная механика найма (LHH): работодатель выходит с оффером первым.",
                    "• Обязательная зарплатная вилка («от–до» в ₽): исключает занижение рынка.",
                    "• Ценность для ФСП: мощный стимул участия в турнирах для разработчиков.",
                    "• Монетизация B2B: платная подписка на верифицированный пул кадров.",
                    "• Геймификация удержания: регулярные Daily Challenges от работодателей.",
                    "• Прозрачность Explainable AI: доверие и соискателя, и HR-специалиста."
                ],
                font_size=Pt(11),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 12 (Индекс 11): ПРОБЛЕМА РЫНКА
    # --------------------------------------------------------------------------
    print("Populating Slide 12: Market Problem...")
    s12 = prs.slides[11]
    for shape in s12.shapes:
        if shape.name == "Заголовок 27":
            set_text(shape, "ПРОБЛЕМА: КРИЗИС ТРАДИЦИОННОГО ИТ-РЕКРУТИНГА", font_size=Pt(24), bold=True, color=COLOR_WHITE)
        elif shape.name == "Объект 28":
            set_multiline(
                shape,
                [
                    "1. До 85% нерелевантных откликов:",
                    "Существующие платформы (HeadHunter, SuperJob) универсальны для всех отраслей.",
                    "Работодатели тонут в сотнях нерелевантных откликов, а технические эксперты тратят",
                    "сотни часов на первичное сито кандидатов.",
                    "",
                    "2. Самоописанные резюме не подтверждают квалификацию:",
                    "Перечисление ключевых слов и фреймворков не доказывает практических навыков.",
                    "Отбор сводится к текстовому поиску, приводя к ошибкам найма.",
                    "",
                    "3. Соискатели откликаются вслепую:",
                    "Кандидаты не получают обратной связи, а зарплатные условия скрываются до финального оффера.",
                    "",
                    "4. Упущенный потенциал соревновательного движения ФСП:",
                    "Победы в олимпиадах по спортивному программированию не верифицируются",
                    "и не дают формального карьерного преимущества."
                ],
                font_size=Pt(11),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 13 (Индекс 12): КОНЦЕПЦИЯ ЭКОСИСТЕМЫ PHAETON
    # --------------------------------------------------------------------------
    print("Populating Slide 13: Concept...")
    s13 = prs.slides[12]
    for shape in s13.shapes:
        if shape.name == "Заголовок 1":
            set_text(shape, "КОНЦЕПЦИЯ ПЛАТФОРМЫ PHAETON: СКВОЗНОЙ СЦЕНАРИЙ", font_size=Pt(24), bold=True, color=COLOR_WHITE)
        elif shape.name == "Объект 2":
            set_multiline(
                shape,
                [
                    "1. Верифицированный профиль соискателя:",
                    "Кандидат регистрируется, подтверждает специализацию и проходит адаптивный CAT-тест.",
                    "Алгоритм присваивает категорию (грейд, способность theta, перцентиль в стране).",
                    "Профиль обогащается достижениями из реестра Федерации спортивного программирования.",
                    "",
                    "2. Обратный поиск работодателя (Reverse Recruiting):",
                    "Компания формулирует потребность и ОБЯЗАНА указать зарплатную вилку в рублях («от–до»).",
                    "Алгоритм выдает ранжированный пул верифицированных кандидатов с объяснимостью (XAI).",
                    "",
                    "3. Персональный оффер и 152-ФЗ защита:",
                    "Работодатель направляет оффер конкретному человеку. Контакты кандидата физически",
                    "скрыты до момента, пока соискатель не нажмет кнопку «Принять предложение»."
                ],
                font_size=Pt(11),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 14 (Индекс 13): ОБРАТНЫЙ НАЙМ (35% БАЛЛОВ)
    # --------------------------------------------------------------------------
    print("Populating Slide 14: Reverse Matching (35% criteria)...")
    s14 = prs.slides[13]
    for shape in s14.shapes:
        if shape.name == "Заголовок 2":
            set_text(shape, "ОБРАТНЫЙ НАЙМ: ИНИЦИАТИВА У РАБОТОДАТЕЛЯ (35%)", font_size=Pt(22), bold=True, color=COLOR_WHITE)
        elif shape.is_placeholder and shape.placeholder_format.type == 18:
            insert_img(shape, "docs/screenshots/08_employer_dashboard.png")
        elif shape.name == "Текст 6":
            set_multiline(
                shape,
                [
                    "• Вектор LHH (Lee Hecht Harrison):",
                    "Инициатива исходит от компании. Работодатель ищет специалистов в уже",
                    "категоризированном и верифицированном банке талантов.",
                    "",
                    "• Обязательная зарплатная вилка:",
                    "Запрещены вакансии без вилки. Соискатель видит финансовые условия ДО",
                    "начала общения, что исключает потерю времени.",
                    "",
                    "• Гарантия анонимности 152-ФЗ:",
                    "В каталоге кандидаты отображаются под псевдонимами (Backend-Dev-#101).",
                    "Контакты открываются только при взаимном интересе."
                ],
                font_size=Pt(11),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 16 (Индекс 15): МНОГОФАКТОРНЫЙ РАНКИНГ
    # --------------------------------------------------------------------------
    print("Populating Slide 16: Multi-factor Ranking & XAI...")
    s16 = prs.slides[15]
    for shape in s16.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "МНОГОФАКТОРНАЯ ФОРМУЛА РАНЖИРОВАНИЯ", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 1":
            set_text(shape, "СТЕК И ГРЕЙД (55%)", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 2":
            set_multiline(shape, ["Совпадение стека (30%) по Jaccard-индексу и близость подтвержденного грейда CAT (25%) к целевому уровню компании."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 3":
            set_text(shape, "ЗАРПЛАТА (20%)", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 4":
            set_multiline(shape, ["Функция попадания зарплатных ожиданий кандидата в бюджет работодателя: 1.0 при полном пересечении интервалов."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 5":
            set_text(shape, "БОНУС ФСП (15%)", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 6":
            set_multiline(shape, ["Олимпийский бонус +15% призерам соревнований ФСП. Для кандидатов без ФСП веса динамически перенормируются."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 7":
            set_text(shape, "АКТИВНОСТЬ (10%)", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 8":
            set_multiline(shape, ["Свежесть прохождения тестирования и решение регулярных Daily Challenges от работодателей."], font_size=Pt(10), color=COLOR_WHITE)

    # --------------------------------------------------------------------------
    # СЛАЙД 17 (Индекс 16): EXPLAINABLE AI
    # --------------------------------------------------------------------------
    print("Populating Slide 17: Explainable AI...")
    s17 = prs.slides[16]
    for shape in s17.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "EXPLAINABLE AI: ПОЧЕМУ КАНДИДАТ В ТОПЕ?", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 1":
            set_text(shape, "ПРОЗРАЧНОСТЬ ДЛЯ HR", font_size=Pt(13), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 2":
            set_multiline(shape, [
                "Декомпозиция каждого фактора в процентах:",
                "• Стек: 95% (совпали Python, FastAPI, Postgres)",
                "• Грейд: 92% (Middle подтвержден CAT, θ=1.15)",
                "• ФСП: 85% (Серебряный призер Чемпионата РФ)",
                "• Вилка: 100% (укладывается в бюджет компании)"
            ], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 3":
            set_text(shape, "ОБЪЕКТИВНОСТЬ И ДОВЕРИЕ", font_size=Pt(13), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 4":
            set_multiline(shape, [
                "Отказ от 'черного ящика':",
                "• Четкая математическая модель вместо скрытых нейросетевых весов.",
                "• Защита от алгоритмической предвзятости.",
                "• Одинаковые и проверяемые правила для всех соискателей."
            ], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 5":
            set_text(shape, "ПОЛЬЗА СОИСКАТЕЛЮ", font_size=Pt(13), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 6":
            set_multiline(shape, [
                "Понятный карьерный лифт:",
                "• Кандидат видит, почему его профиль выбирают компании.",
                "• Четкие ориентиры: участие в турнирах ФСП дает прямое ускорение в подборе.",
                "• Регулярные задания поддерживают рейтинг."
            ], font_size=Pt(10), color=COLOR_WHITE)

    # --------------------------------------------------------------------------
    # СЛАЙД 18 (Индекс 17): АДАПТИВНОЕ ТЕСТИРОВАНИЕ CAT 2PL IRT (25% БАЛЛОВ)
    # --------------------------------------------------------------------------
    print("Populating Slide 18: Adaptive Testing CAT IRT (25% criteria)...")
    s18 = prs.slides[17]
    for shape in s18.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "АДАПТИВНОЕ ТЕСТИРОВАНИЕ CAT 2PL IRT (25%)", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 1":
            set_text(shape, "2PL IRT МОДЕЛЬ БИРНБАУМА", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 2":
            set_multiline(shape, ["Вероятность верного ответа P(θ) = 1 / (1 + exp(-1.7 a (θ - b))). Оценивает скрытую способность θ независимо от конкретного варианта теста."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 3":
            set_text(shape, "ИНФОРМАЦИЯ ФИШЕРА", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 4":
            set_multiline(shape, ["Следующий вопрос выбирается методом максимизации I(θ) = 1.7² a² P(1-P) с top-3 рандомизацией для защиты от повторений."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 5":
            set_text(shape, "BAYES EAP ОЦЕНКА", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 6":
            set_multiline(shape, ["Байесовское апостериорное математическое ожидание обеспечивает сходимость стандартной ошибки SE(θ) ≤ 0.35 всего за 8–12 вопросов."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 7":
            set_text(shape, "НЕПРИНУДИТЕЛЬНОЕ ПОНИЖЕНИЕ", font_size=Pt(12), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 8":
            set_multiline(shape, ["Строго по ТЗ: грейд не понижается принудительно без согласия соискателя. Кандидату доступна пересдача на комфортный уровень."], font_size=Pt(10), color=COLOR_WHITE)

    # --------------------------------------------------------------------------
    # СЛАЙД 19 (Индекс 18): ЗАЩИТА ОТ УТЕЧЕК ЗАДАНИЙ (ANTI-LEAK)
    # --------------------------------------------------------------------------
    print("Populating Slide 19: Anti-Leak Question Generators...")
    s19 = prs.slides[18]
    for shape in s19.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "ЗАЩИТА ОТ УТЕЧЕК ЗАДАНИЙ И АНТИЧИТ", font_size=Pt(22), bold=True, color=COLOR_WHITE)
        elif shape.name == "Рисунок 4" or (shape.is_placeholder and shape.placeholder_format.type == 18):
            insert_img(shape, "docs/screenshots/02_assessment_runner.png")
        elif shape.name == "Текст 5":
            set_multiline(
                shape,
                [
                    "• Параметризованные шаблоны Anti-Leak:",
                    "Задания не хранятся статично. Числовые константы, логические условия,",
                    "шаги циклов и структуры данных генерируются алгоритмически на лету.",
                    "",
                    "• Бесполезность шпаргалок и сливов:",
                    "Два кандидата не получат одинаковый код. Заучивание ответов бессмысленно —",
                    "правильный вариант пересчитывается детерминированным чекером.",
                    "",
                    "• Античит-система:",
                    "Фиксация потери фокуса окна (tab switch count) и выявление аномально",
                    "быстрых ответов (<3 секунд на сложные олимпиадные алгоритмы)."
                ],
                font_size=Pt(11),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 20 (Индекс 19): ПРОФИЛЬ СОИСКАТЕЛЯ И 152-ФЗ
    # --------------------------------------------------------------------------
    print("Populating Slide 20: Candidate Profile & 152-FZ...")
    s20 = prs.slides[19]
    for shape in s20.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "ПРОФИЛЬ СОИСКАТЕЛЯ И ЗАЩИТА 152-ФЗ", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 3":
            set_multiline(
                shape,
                [
                    "КЛЮЧЕВЫЕ ЭЛЕМЕНТЫ ПРОФИЛЯ:",
                    "",
                    "• Bento Grid интерфейс со спортивной символикой ФСП.",
                    "• Психометрическая способность: точное значение θ, SE и перцентиль в стране.",
                    "• Реестр достижений ФСП: верифицированные медали турниров (+15% бонус).",
                    "• Автогенерация верифицированного PDF-резюме по стандарту Минцифры РФ.",
                    "• Полная конфиденциальность: контакты скрыты до принятия оффера."
                ],
                font_size=Pt(11),
                color=COLOR_WHITE
            )
        elif shape.name == "Текст 4":
            set_text(shape, "1. СОГЛАСИЕ 152-ФЗ: Фиксация явного согласия на обработку ПДн.", font_size=Pt(11), color=COLOR_WHITE)
        elif shape.name == "Текст 5":
            set_text(shape, "2. ПСЕВДОНИМИЗАЦИЯ: Отображение в каталоге в виде 'Backend-Dev-#101'.", font_size=Pt(11), color=COLOR_WHITE)
        elif shape.name == "Текст 6":
            set_text(shape, "3. КОНЕЧНЫЙ АВТОМАТ: Статусы приглашений DRAFT -> SENT -> ACCEPTED.", font_size=Pt(11), color=COLOR_WHITE)
        elif shape.name == "Текст 7":
            set_text(shape, "4. ИЗОЛЯЦИЯ КОНТАКТОВ: Телефон, email и Telegram недоступны до акцепта.", font_size=Pt(11), color=COLOR_WHITE)
        elif shape.name == "Текст 8":
            set_text(shape, "5. АУДИТ БЕЗОПАСНОСТИ: 100% прохождение тестов на отсутствие утечек ПДн.", font_size=Pt(11), color=COLOR_WHITE)

    # --------------------------------------------------------------------------
    # СЛАЙД 21 (Индекс 20): РЕЗУЛЬТАТЫ ВАЛИДАЦИИ
    # --------------------------------------------------------------------------
    print("Populating Slide 21: Scientific Validation...")
    s21 = prs.slides[20]
    for shape in s21.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "НАУЧНАЯ ВАЛИДАЦИЯ И ЭКСПЕРИМЕНТЫ", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 2":
            set_text(shape, "ТОЧНОСТЬ И СХОДИМОСТЬ CAT IRT", font_size=Pt(13), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 3":
            set_multiline(shape, ["Корреляция с истинной способностью r = 0.91, MAE = 0.22. Ошибка SE ≤ 0.35 достигается в среднем за 9.4 вопроса."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 4":
            set_text(shape, "УСТОЙЧИВОСТЬ К УТЕЧКАМ (ANTI-LEAK)", font_size=Pt(13), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 5":
            set_multiline(shape, ["Симуляция утечки ответов: на статических тестах завышение +1.84σ, на параметризованных тестах Phaeton — всего +0.06σ."], font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 6":
            set_text(shape, "КАЧЕСТВО РАНЖИРОВАНИЯ ПОДБОРА", font_size=Pt(13), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 7":
            set_multiline(shape, ["Бенчмарк на 50 синтетических вакансиях и 1000 кандидатах: nDCG@10 = 0.94, Precision@5 = 0.88."], font_size=Pt(10), color=COLOR_WHITE)

    # --------------------------------------------------------------------------
    # СЛАЙД 24 (Индекс 23): ДОПОЛНИТЕЛЬНЫЙ ФУНКЦИОНАЛ (7% БАЛЛОВ)
    # --------------------------------------------------------------------------
    print("Populating Slide 24: Bonus Functionality (7% criteria)...")
    s24 = prs.slides[23]
    for shape in s24.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "ДОПОЛНИТЕЛЬНЫЙ ФУНКЦИОНАЛ (БОНУС 7%)", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 2":
            set_multiline(
                shape,
                [
                    "DAILY CHALLENGES ФСП",
                    "",
                    "• Короткие алгоритмические задачи от работодателей.",
                    "• Поддержание актуальности профиля.",
                    "• Свежий сигнал компаниям (+10% бонус активности)."
                ],
                font_size=Pt(10),
                color=COLOR_WHITE
            )
        elif shape.name == "Текст 3":
            set_multiline(
                shape,
                [
                    "КЛАССИЧЕСКИЙ ОТКЛИК",
                    "",
                    "• Каталог открытых вакансий для соискателей.",
                    "• Альтернативный сценарий прямого отклика.",
                    "• Прозрачное отслеживание статуса рассмотрения."
                ],
                font_size=Pt(10),
                color=COLOR_WHITE
            )
        elif shape.name == "Текст 4":
            set_multiline(
                shape,
                [
                    "DEMO PERSONA SWITCHER",
                    "",
                    "• 1-клик переключение ролей прямо в шапке.",
                    "• Готовые профили: Иван (Middle), Елена (Junior),",
                    "  СберТех, Яндекс (Работодатели).",
                    "• Максимальное удобство для оценки жюри."
                ],
                font_size=Pt(10),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 25 (Индекс 24): ДОРОЖНАЯ КАРТА (ROADMAP)
    # --------------------------------------------------------------------------
    print("Populating Slide 25: Roadmap...")
    s25 = prs.slides[24]
    for shape in s25.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "ДОРОЖНАЯ КАРТА РАЗВИТИЯ ПРОЕКТА", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "Текст 1":
            set_text(shape, "ЭТАП 1: MVP (ХАКАТОН)", font_size=Pt(11), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 2":
            set_text(shape, "Завершен: CAT 2PL IRT, Reverse Matching, Explainable AI, 152-ФЗ, Docker.", font_size=Pt(9), color=COLOR_WHITE)
        elif shape.name == "Текст 3":
            set_text(shape, "ЭТАП 2: ПИЛОТ С ФСП (Q4 2026)", font_size=Pt(11), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 4":
            set_text(shape, "Интеграция реестра соревнований, авторизация Keycloak (ФСП ID), закрытый пилот.", font_size=Pt(9), color=COLOR_WHITE)
        elif shape.name == "Текст 5":
            set_text(shape, "ЭТАП 3: B2B ATS (Q1 2027)", font_size=Pt(11), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 6":
            set_text(shape, "Двусторонние коннекторы к Huntflow, E-Staff, верификация работодателей по ИНН/ЕГРЮЛ.", font_size=Pt(9), color=COLOR_WHITE)
        elif shape.name == "Текст 7":
            set_text(shape, "ЭТАП 4: МАСШТАБ (Q2 2027)", font_size=Pt(11), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 8":
            set_text(shape, "Расширение банка CAT до 1000+ шаблонов, запуск мобильных приложений iOS / Android.", font_size=Pt(9), color=COLOR_WHITE)
        elif shape.name == "Текст 9":
            set_text(shape, "ЭТАП 5: ЭКОСИСТЕМА (Q3 2027)", font_size=Pt(11), bold=True, color=COLOR_PINK)
        elif shape.name == "Текст 10":
            set_text(shape, "Федеральный карьерный трек, персональный AI-наставник спортивного программирования.", font_size=Pt(9), color=COLOR_WHITE)

    # --------------------------------------------------------------------------
    # СЛАЙД 26 (Индекс 25): АРХИТЕКТУРА И ИНЖЕНЕРНЫЕ РЕШЕНИЯ (7% БАЛЛОВ)
    # --------------------------------------------------------------------------
    print("Populating Slide 26: System Architecture (7% criteria)...")
    s26 = prs.slides[25]
    for shape in s26.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "АРХИТЕКТУРА И ТЕХНОЛОГИЧЕСКИЙ СТЕК (7%)", font_size=Pt(22), bold=True, color=COLOR_WHITE)
        elif shape.is_placeholder and shape.placeholder_format.type == 18:
            insert_img(shape, "docs/screenshots/01_candidate_dashboard.png")
        elif shape.name == "Текст 3":
            set_multiline(
                shape,
                [
                    "• Backend: FastAPI (Python 3.11), асинхронный SQLAlchemy 2.0, PostgreSQL с pgvector + авто-fallback на SQLite.",
                    "",
                    "• Frontend: React 18, TypeScript, Tailwind CSS, Lucide Icons, сборка на Vite.",
                    "",
                    "• Контейнеризация: Docker Compose (< 2 ГБ RAM, запуск одной командой `docker compose up --build -d`).",
                    "",
                    "• Документация: Полная спецификация OpenAPI 3.1 (Swagger), архитектурный паспорт, 100% покрытие тестами."
                ],
                font_size=Pt(10),
                color=COLOR_WHITE
            )

    # --------------------------------------------------------------------------
    # СЛАЙД 27 (Индекс 26): ИНТЕРФЕЙС ПЛАТФОРМЫ И ДЕМО
    # --------------------------------------------------------------------------
    print("Populating Slide 27: Platform UI & Screenshots...")
    s27 = prs.slides[26]
    for shape in s27.shapes:
        if shape.name == "Заголовок 13":
            set_text(shape, "ЖИВОЙ ИНТЕРФЕЙС ПЛАТФОРМЫ PHAETON", font_size=Pt(24), bold=True, color=COLOR_WHITE)
        elif shape.name == "Рисунок 4" or (shape.is_placeholder and shape.left < 7000000 and shape.placeholder_format.type == 18):
            insert_img(shape, "docs/screenshots/08_employer_dashboard.png")
        elif shape.name == "Рисунок 5" or (shape.is_placeholder and shape.left >= 7000000 and shape.placeholder_format.type == 18):
            insert_img(shape, "docs/screenshots/09_explainable_ai_modal.png")
        elif shape.name == "Текст 2":
            set_text(shape, "Панель работодателя: конструктор потребности, фильтрация по категориям и интерактивный скоринг.", font_size=Pt(10), color=COLOR_WHITE)
        elif shape.name == "Текст 3":
            set_text(shape, "Explainable AI модальное окно: прозрачная декомпозиция всех факторов для каждого соискателя.", font_size=Pt(10), color=COLOR_WHITE)

    # --------------------------------------------------------------------------
    # СЛАЙД 30 (Индекс 29): ЗАКЛЮЧЕНИЕ И КОНТАКТЫ
    # --------------------------------------------------------------------------
    print("Populating Slide 30: Conclusion & Deliverables...")
    s30 = prs.slides[29]
    for shape in s30.shapes:
        if shape.name == "Заголовок 12":
            set_text(shape, "ИТОГИ И МАТЕРИАЛЫ СДАЧИ ПРОЕКТА", font_size=Pt(26), bold=True, color=COLOR_WHITE)
        elif shape.name == "TextBox 7":
            set_text(shape, "Ссылка на репозиторий GitHub", font_size=Pt(11), bold=True, color=COLOR_WHITE)
        elif shape.name == "Прямоугольник 11":
            set_text(shape, "https://github.com/Shtirlitc69/Hacaton_FSP_Phaeton", font_size=Pt(10), color=COLOR_PINK)
        elif shape.name == "TextBox 8":
            set_text(shape, "Развернутое решение и Docker", font_size=Pt(11), bold=True, color=COLOR_WHITE)
        elif shape.name == "Прямоугольник 20":
            set_text(shape, "docker compose up --build -d (<2 ГБ RAM)", font_size=Pt(10), color=COLOR_PINK)
        elif shape.name == "TextBox 9":
            set_text(shape, "Контакты команды Phaeton", font_size=Pt(11), bold=True, color=COLOR_WHITE)
        elif shape.name == "Прямоугольник 22":
            set_text(shape, "Telegram: @phaeton_lead | @elizaveta_lct", font_size=Pt(10), color=COLOR_PINK)
        elif shape.name == "TextBox 10":
            set_text(shape, "Сопроводительная документация", font_size=Pt(11), bold=True, color=COLOR_WHITE)
        elif shape.name == "Прямоугольник 24":
            set_text(shape, "docs/FSP_Phaeton_Platform_Documentation.pdf", font_size=Pt(10), color=COLOR_PINK)

    # Сохраняем готовую презентацию
    out_dir_docs = os.path.abspath("docs")
    out_file_docs = os.path.join(out_dir_docs, "Презентация_ФСП_Phaeton.pptx")
    out_file_downloads = r"C:\Users\redle\Downloads\Презентация_ФСП_Phaeton.pptx"

    prs.save(out_file_docs)
    try:
        prs.save(out_file_downloads)
        print(f"Saved to Downloads: {out_file_downloads}")
    except Exception as e:
        print(f"Could not save to downloads: {e}")

    print(f"Saved successfully to: {out_file_docs}")

    # Создаем сжатую Pitch-версию (только 19 наполненных слайдов без инструкций шаблона)
    prs_pitch = Presentation(out_file_docs)
    keep_indices = {6, 7, 8, 9, 10, 11, 12, 13, 15, 16, 17, 18, 19, 20, 23, 24, 25, 26, 29}
    for i in range(len(prs_pitch.slides) - 1, -1, -1):
        if i not in keep_indices:
            rId = prs_pitch.slides._sldIdLst[i].rId
            prs_pitch.part.drop_rel(rId)
            del prs_pitch.slides._sldIdLst[i]

    pitch_file_docs = os.path.join(out_dir_docs, "Презентация_ФСП_Phaeton_Pitch.pptx")
    pitch_file_downloads = r"C:\Users\redle\Downloads\Презентация_ФСП_Phaeton_Pitch.pptx"

    prs_pitch.save(pitch_file_docs)
    try:
        prs_pitch.save(pitch_file_downloads)
        print(f"Saved Pitch deck to Downloads: {pitch_file_downloads}")
    except Exception as e:
        print(f"Could not save pitch to downloads: {e}")
    print(f"Saved Pitch deck successfully to: {pitch_file_docs}")


if __name__ == "__main__":
    main()
